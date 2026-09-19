"""회장·정원 통합 마스터 표의 파일 입출력 — CSV / Excel.

회장관리(bulk-venues)와 정원관리(bulk-capacity) 모두 동일한 표준 양식으로 내보낸다.
원본 마스터(data/hospitals/*.xlsx)와 동일한 32개 열 구조를 갖추며:
- 회장 소재지/시설 정보 (ID, 会場番号, 地域, 区・市町村, 会場名, 住所, 開始時刻, 受付終了, 駐車場, 緯度, 経度, 交通情報, アクセス時間)
- 개최 회차 정보 (開催日, 予約終了)
- 16개 시간대별 정원 (9:00～09:30, 9:30～10:00, ..., 16:30～17:00)
- 합계 (計)
를 순서대로 출력한다.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime, time
from typing import Any

from sqlalchemy.orm import Session

from app.core import time_grid
from app.models.hospital import Hospital, HospitalSchedule
from app.services.sheet_io import SheetError, _openpyxl

# 32개 표준 열 정의: (키, 표시 라벨, Excel 열 폭)
MASTER_COLUMNS: list[tuple[str, str, int]] = [
    ("id", "ID", 6),
    ("venue_no", "会場番号", 12),
    ("area", "地域", 14),
    ("city", "区・市町村", 16),
    ("name", "会場名", 30),
    ("address", "住所", 38),
    ("open_time", "開始時刻", 12),
    ("reception_end_time", "受付終了", 12),
    ("has_parking", "駐車場", 10),
    ("latitude", "緯度", 14),
    ("longitude", "経度", 14),
    ("transit_info", "交通情報", 34),
    ("access_minutes", "アクセス時間", 14),
    ("event_date", "開催日", 14),
    ("booking_close_date", "予約終了", 14),
    ("slot_0900", "9:00～09:30", 12),
    ("slot_0930", "9:30～10:00", 12),
    ("slot_1000", "10:00～10:30", 12),
    ("slot_1030", "10:30～11:00", 12),
    ("slot_1100", "11:00～11:30", 12),
    ("slot_1130", "11:30～12:00", 12),
    ("slot_1200", "12:00～12:30", 12),
    ("slot_1230", "12:30～13:00", 12),
    ("slot_1300", "13:00～13:30", 12),
    ("slot_1330", "13:30～14:00", 12),
    ("slot_1400", "14:00～14:30", 12),
    ("slot_1430", "14:30～15:00", 12),
    ("slot_1500", "15:00～15:30", 12),
    ("slot_1530", "15:30～16:00", 12),
    ("slot_1600", "16:00～16:30", 12),
    ("slot_1630", "16:30～17:00", 12),
    ("total", "計", 10),
]

# 시간대 슬롯 인덱스별 매핑
SLOT_KEYS = [
    (i, time_grid.capacity_column(i), MASTER_COLUMNS[15 + i][1])
    for i in range(len(time_grid.TIME_GRID))
]


def _format_time(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, (datetime, time)):
        return val.strftime("%H:%M")
    text = str(val).strip()
    m = re.match(r"^(\d{1,2}):(\d{2})", text)
    if m:
        return f"{int(m.group(1)):02d}:{int(m.group(2)):02d}"
    return text


def _format_date(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    text = str(val).strip()
    m = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", text)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    return text


def _format_parking(val: Any) -> str:
    if val is True or val in ("○", "◯", "はい", "1", 1, "true", "True", "有", "あり"):
        return "○"
    return "×"


def build_master_rows(
    db: Session,
    sheet: str,
    payload_rows: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    """회장관리(venues) 또는 정원관리(capacity)의 데이터를 통일된 마스터 행 목록으로 생성한다."""
    payload_rows = payload_rows or []

    # 전체 회장 및 회차 맵 로드
    hospitals = (
        db.query(Hospital)
        .order_by(Hospital.sort_order, Hospital.id)
        .all()
    )
    h_by_code: dict[str, Hospital] = {}
    h_by_id: dict[str, Hospital] = {}
    for h in hospitals:
        h_by_code[h.code] = h
        if h.code.startswith("V") and h.code[1:].isdigit():
            h_by_code[str(int(h.code[1:]))] = h
        h_by_id[str(h.id)] = h

    rows: list[dict[str, Any]] = []

    # 1. 정원관리(capacity)에서 내보내기를 호출한 경우
    if sheet == "capacity" and payload_rows:
        for idx, p_row in enumerate(payload_rows, 1):
            code = str(p_row.get("hospital_code") or p_row.get("code") or "")
            h = h_by_code.get(code) or h_by_id.get(str(p_row.get("hospital_id", "")))

            # 회장 기본 정보
            venue_no = code or (h.code if h else "")
            area = (h.area if h else "") or p_row.get("area", "")
            city = (h.city if h else "") or p_row.get("city", "")
            name = p_row.get("hospital_name") or (h.name if h else "")
            address = (h.address if h else "") or p_row.get("address", "")
            has_parking = _format_parking(h.has_parking if h else p_row.get("has_parking"))
            latitude = str(h.latitude if (h and h.latitude is not None) else (p_row.get("latitude") or ""))
            longitude = str(h.longitude if (h and h.longitude is not None) else (p_row.get("longitude") or ""))
            transit_info = (h.transit_info if h else "") or p_row.get("transit_info", "")
            access_minutes = (h.access_minutes if h else "") or p_row.get("access_minutes", "")

            # 회차 정보
            open_time = _format_time(p_row.get("open_time"))
            reception_end_time = _format_time(p_row.get("reception_end_time"))
            event_date = _format_date(p_row.get("event_date"))
            booking_close_date = _format_date(p_row.get("booking_close_date"))

            # 슬롯 정원 값 추출 및 합계 계산
            slot_values: dict[str, Any] = {}
            total = 0
            for s_idx, cap_col, s_label in SLOT_KEYS:
                # p_row에서 다양한 키 형태 지원 (예: '09:00～09:30', '9:00～09:30', '9:00~09:30', 'cap_0900')
                alt1 = time_grid.label(s_idx)  # '09:00～09:30'
                alt2 = s_label                # '9:00～09:30'
                alt3 = alt1.replace("～", "~")
                alt4 = alt2.replace("～", "~")

                raw_val = (
                    p_row.get(alt1)
                    if alt1 in p_row
                    else p_row.get(alt2)
                    if alt2 in p_row
                    else p_row.get(alt3)
                    if alt3 in p_row
                    else p_row.get(alt4)
                    if alt4 in p_row
                    else p_row.get(cap_col)
                )

                if raw_val is not None and str(raw_val).strip() != "":
                    try:
                        num = int(str(raw_val).strip())
                        slot_values[s_label] = num
                        total += num
                    except ValueError:
                        slot_values[s_label] = str(raw_val).strip()
                else:
                    slot_values[s_label] = ""

            row_dict: dict[str, Any] = {
                "ID": idx,
                "会場番号": venue_no,
                "地域": area,
                "区・市町村": city,
                "会場名": name,
                "住所": address,
                "開始時刻": open_time,
                "受付終了": reception_end_time,
                "駐車場": has_parking,
                "緯度": latitude,
                "経度": longitude,
                "交通情報": transit_info,
                "アクセス時間": access_minutes,
                "開催日": event_date,
                "予約終了": booking_close_date,
            }
            row_dict.update(slot_values)
            row_dict["計"] = total if total > 0 else ""
            rows.append(row_dict)

    # 2. 회장관리(venues)에서 내보내기를 호출한 경우
    elif sheet == "venues" and payload_rows:
        row_counter = 1
        for p_row in payload_rows:
            code = str(p_row.get("code") or "")
            h = h_by_code.get(code) or h_by_id.get(str(p_row.get("id", "")))

            # 화면에서 수정한 최신 회장 정보 우선 반영
            venue_no = code or (h.code if h else "")
            area = p_row.get("area") or (h.area if h else "")
            city = p_row.get("city") or (h.city if h else "")
            name = p_row.get("name") or (h.name if h else "")
            address = p_row.get("address") or (h.address if h else "")
            has_parking = _format_parking(p_row.get("has_parking") if "has_parking" in p_row else (h.has_parking if h else None))
            latitude = str(p_row.get("latitude") if p_row.get("latitude") is not None else (h.latitude if (h and h.latitude is not None) else ""))
            longitude = str(p_row.get("longitude") if p_row.get("longitude") is not None else (h.longitude if (h and h.longitude is not None) else ""))
            transit_info = p_row.get("transit_info") or (h.transit_info if h else "")
            access_minutes = p_row.get("access_minutes") or (h.access_minutes if h else "")

            # 해당 회장의 회차 목록
            schedules = sorted(h.schedules, key=lambda s: s.event_date) if (h and h.schedules) else []

            if schedules:
                for s in schedules:
                    open_time = _format_time(s.open_time)
                    reception_end_time = _format_time(s.reception_end_time)
                    event_date = _format_date(s.event_date)
                    booking_close_date = _format_date(s.booking_close_date)

                    slot_values = {}
                    total = 0
                    for s_idx, cap_col, s_label in SLOT_KEYS:
                        val = getattr(s, cap_col, None)
                        if val is not None:
                            slot_values[s_label] = val
                            total += val
                        else:
                            slot_values[s_label] = ""

                    row_dict = {
                        "ID": row_counter,
                        "会場番号": venue_no,
                        "地域": area,
                        "区・市町村": city,
                        "会場名": name,
                        "住所": address,
                        "開始時刻": open_time,
                        "受付終了": reception_end_time,
                        "駐車場": has_parking,
                        "緯度": latitude,
                        "経度": longitude,
                        "交通情報": transit_info,
                        "アクセス時間": access_minutes,
                        "開催日": event_date,
                        "予約終了": booking_close_date,
                    }
                    row_dict.update(slot_values)
                    row_dict["計"] = total if total > 0 else ""
                    rows.append(row_dict)
                    row_counter += 1
            else:
                # 회차가 없는 경우 회장 정보만 1행 출력
                row_dict = {
                    "ID": row_counter,
                    "会場番号": venue_no,
                    "地域": area,
                    "区・市町村": city,
                    "会場名": name,
                    "住所": address,
                    "開始時刻": "",
                    "受付終了": "",
                    "駐車場": has_parking,
                    "緯度": latitude,
                    "経度": longitude,
                    "交通情報": transit_info,
                    "アクセス時間": access_minutes,
                    "開催日": "",
                    "予約終了": "",
                }
                for _, _, s_label in SLOT_KEYS:
                    row_dict[s_label] = ""
                row_dict["計"] = ""
                rows.append(row_dict)
                row_counter += 1

    # 3. 전체 DB 기반 출력 (기본 폴백)
    if not rows:
        all_schedules = (
            db.query(HospitalSchedule, Hospital)
            .join(Hospital, Hospital.id == HospitalSchedule.hospital_id)
            .order_by(Hospital.sort_order, Hospital.id, HospitalSchedule.event_date)
            .all()
        )
        for idx, (s, h) in enumerate(all_schedules, 1):
            slot_values = {}
            total = 0
            for s_idx, cap_col, s_label in SLOT_KEYS:
                val = getattr(s, cap_col, None)
                if val is not None:
                    slot_values[s_label] = val
                    total += val
                else:
                    slot_values[s_label] = ""

            row_dict = {
                "ID": idx,
                "会場番号": h.code,
                "地域": h.area,
                "区・市町村": h.city,
                "会場名": h.name,
                "住所": h.address,
                "開始時刻": _format_time(s.open_time),
                "受付終了": _format_time(s.reception_end_time),
                "駐車場": _format_parking(h.has_parking),
                "緯度": str(h.latitude) if h.latitude is not None else "",
                "経度": str(h.longitude) if h.longitude is not None else "",
                "交通情報": h.transit_info or "",
                "アクセス時間": h.access_minutes or "",
                "開催日": _format_date(s.event_date),
                "予約終了": _format_date(s.booking_close_date),
            }
            row_dict.update(slot_values)
            row_dict["計"] = total if total > 0 else ""
            rows.append(row_dict)

    return rows


def build_csv(rows: list[dict[str, Any]]) -> bytes:
    """UTF-8 BOM + CRLF 형식의 CSV를 생성한다."""
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
    headers = [label for _, label, _ in MASTER_COLUMNS]
    writer.writerow(headers)

    for row in rows:
        writer.writerow([row.get(label, "") for label in headers])

    return buffer.getvalue().encode("utf-8-sig")


def build_xlsx(rows: list[dict[str, Any]]) -> bytes:
    """모든 셀을 텍스트 및 숫자로 깔끔하게 기록한 표준 Excel(.xlsx) 파일을 생성한다."""
    Workbook, _ = _openpyxl()
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = "会場マスター"

    headers = [label for _, label, _ in MASTER_COLUMNS]
    sheet.append(headers)

    header_font = Font(name="Yu Gothic", size=10, bold=True, color="FF1F3B34")
    header_fill = PatternFill("solid", fgColor="FFE6EFED")
    header_align = Alignment(vertical="center", horizontal="center", wrap_text=False)
    thin_border = Border(
        left=Side(style="thin", color="FFD1D5DB"),
        right=Side(style="thin", color="FFD1D5DB"),
        top=Side(style="thin", color="FFD1D5DB"),
        bottom=Side(style="thin", color="FFD1D5DB"),
    )

    for col_idx, (_, _, width) in enumerate(MASTER_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        sheet.column_dimensions[get_column_letter(col_idx)].width = width

    sheet.row_dimensions[1].height = 24

    data_font = Font(name="Yu Gothic", size=10, color="FF1F2937")
    center_align = Alignment(vertical="center", horizontal="center")
    left_align = Alignment(vertical="center", horizontal="left")
    right_align = Alignment(vertical="center", horizontal="right")

    for r_idx, row in enumerate(rows, start=2):
        sheet.row_dimensions[r_idx].height = 20
        for c_idx, (_, label, _) in enumerate(MASTER_COLUMNS, start=1):
            cell = sheet.cell(row=r_idx, column=c_idx)
            val = row.get(label, "")
            cell.value = val
            cell.font = data_font
            cell.border = thin_border

            # 정렬 서식
            if label in ("ID", "会場番号", "地域", "区・市町村", "開始時刻", "受付終了", "駐車場", "開催日", "予約終了"):
                cell.alignment = center_align
            elif label in ("会場名", "住所", "交通情報", "アクセス時間"):
                cell.alignment = left_align
            else:
                cell.alignment = right_align

    # 틀 고정: 1행 고정
    sheet.freeze_panes = "A2"

    # 자동 필터
    max_col_letter = get_column_letter(len(MASTER_COLUMNS))
    max_row = max(len(rows) + 1, 2)
    sheet.auto_filter.ref = f"A1:{max_col_letter}{max_row}"

    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def export_file_name(kind: str, moment: datetime | None = None) -> str:
    moment = moment or datetime.now()
    return f"会場マスター_一括管理_{moment.strftime('%Y%m%d_%H%M%S')}.{kind}"

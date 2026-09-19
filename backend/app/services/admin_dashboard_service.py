"""관리 화면 — 대시보드 (A-01).

아침에 로그인해서 **가장 먼저 봐야 할 것만** 올린다.
숫자를 늘어놓는 화면이 아니라 「오늘 무엇을 해야 하는가」를 알려 주는 화면이다.

  · 대응이 필요한 예약(임시·불비·휴진)이 몇 건인가      (자체 피드백 M-4)
  · 오늘 검진은 회장별로 몇 명인가
  · 메일 발송이 실패한 건이 있는가                      (자체 피드백 M-12)
  · 곧 자동 삭제될 예약이 몇 건인가
"""

import calendar
import zipfile
import csv
from datetime import date, datetime, timedelta
import io

from sqlalchemy import extract, func, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.exam_option import ExamOption
from app.models.hospital import Hospital, HospitalSchedule
from app.models.admin import AuditLog
from app.models.mail import TEMPLATE_LABELS, MailLog
from app.models.reservation import CANCEL_TYPE_LABELS, Reservation
from app.schemas.admin import (
    DashboardAgeBucket,
    DashboardChannelStats,
    DashboardCounter,
    DashboardData,
    DashboardGenderStats,
    DashboardMailFailure,
    DashboardOptionStat,
    DashboardRecent,
    DashboardSeries,
    DashboardSeriesPoint,
    DashboardStatusStats,
    DashboardTodayRow,
)
from app.services import purge_service
from app.services import slot_grid_service as grid
from app.services.admin_reservation_service import (
    CHANNEL_LABELS,
    STATUS_LABELS,
    is_holiday_slot,
)


_WEEKDAY_JA = ["月", "火", "水", "木", "金", "土", "日"]


def _short_date(d: date) -> str:
    """「2026.8.21 (금)」 — 지표 카드에 들어갈 짧은 날짜.

    연도를 넣는 이유는 화면을 캡처해 공유했을 때 언제 것인지 남기기 위함이다.
    회계연도가 4월에 시작해 연도를 넘나드는 것도 이유다.
    한글 「년·월·일」을 빼면 카드 폭에서 줄바꿈되지 않는다.
    """
    return f"{d.year}.{d.month}.{d.day}（{_WEEKDAY_JA[d.weekday()]}）"


def _recent(r: Reservation) -> DashboardRecent:
    return DashboardRecent(
        id=r.id,
        reservation_no=r.reservation_no,
        full_name=r.full_name,
        hospital_name=r.hospital.name if r.hospital else "",
        slot_date=r.slot_date,
        time_label=r.time_label,
        channel=CHANNEL_LABELS.get(r.channel, r.channel),
        channel_code=r.channel,
        status=STATUS_LABELS.get(r.status, r.status),
        defect_note=r.defect_note or "",
        # 휴대전화를 우선한다. 낮에 연결될 가능성이 높다.
        tel=r.tel_mobile or r.tel_home or "",
        created_at=r.created_at,
    )


def _period_filter(
    stmt,
    *,
    year: int | None = None,
    month: int | None = None,
    date_str: str | None = None,
    week_start: str | None = None,
    week_end: str | None = None,
):
    """통계 화면의 기간(日別·週別·月別)을 申込日(created_at) 기준으로 건다.

    화면 조회와 내보내기가 **같은 조건**을 쓰도록 한 곳에 둔다. 예전에는
    내보내기가 기간을 받지 않아, 화면에서 9/1～9/27 을 골라도 파일에는
    전체 기간 숫자가 나왔다.
    """
    if year:
        stmt = stmt.where(extract("year", Reservation.created_at) == year)
    if month:
        stmt = stmt.where(extract("month", Reservation.created_at) == month)
    if date_str:
        try:
            d_val = datetime.strptime(date_str, "%Y-%m-%d").date()
            stmt = stmt.where(func.date(Reservation.created_at) == d_val)
        except ValueError:
            pass
    if week_start and week_end:
        try:
            ws_val = datetime.strptime(week_start, "%Y-%m-%d").date()
            we_val = datetime.strptime(week_end, "%Y-%m-%d").date()
            stmt = stmt.where(
                func.date(Reservation.created_at) >= ws_val,
                func.date(Reservation.created_at) <= we_val,
            )
        except ValueError:
            pass
    return stmt


def _period_range(
    reservations,
    *,
    year: int | None = None,
    month: int | None = None,
    date_str: str | None = None,
    week_start: str | None = None,
    week_end: str | None = None,
) -> tuple[date | None, date | None]:
    """파일에 적을 기간의 처음과 끝. 조건이 없으면 실제 접수일의 처음과 끝."""
    def parse(value):
        try:
            return datetime.strptime(value, "%Y-%m-%d").date()
        except (TypeError, ValueError):
            return None

    one_day = parse(date_str)
    if one_day:
        return one_day, one_day
    ws_val, we_val = parse(week_start), parse(week_end)
    if ws_val and we_val:
        return ws_val, we_val
    if year and month:
        last = calendar.monthrange(int(year), int(month))[1]
        return date(int(year), int(month), 1), date(int(year), int(month), last)
    if year:
        return date(int(year), 1, 1), date(int(year), 12, 31)
    days = [r.created_at.date() for r in reservations if r.created_at]
    if days:
        return min(days), max(days)
    return None, None


def _stats_series(reservations, start: date | None, end: date | None) -> DashboardSeries:
    """健診統計의 곡선에 쓸 건수.

    하루짜리를 날짜로 그리면 점 하나라서 시간대별로 나눈다. 31일을 넘으면
    날짜가 너무 촘촘해지므로 월별로 묶는다 — CSV 의 日別/月別 규칙과 같다.
    """
    created = [r.created_at for r in reservations if r.created_at]
    if not start or not end:
        return DashboardSeries(kind="day", label="日別", points=[])

    if start == end:
        counts = [0] * 24
        for c in created:
            counts[c.hour] += 1
        return DashboardSeries(
            kind="hour", label="時間帯別", start=start, end=end,
            points=[DashboardSeriesPoint(label=f"{h}時", count=n) for h, n in enumerate(counts)],
        )

    span = (end - start).days + 1
    if span <= 31:
        by_day: dict[date, int] = {}
        for c in created:
            by_day[c.date()] = by_day.get(c.date(), 0) + 1
        days = [start + timedelta(days=i) for i in range(span)]
        return DashboardSeries(
            kind="day", label="日別", start=start, end=end,
            points=[DashboardSeriesPoint(label=f"{d.month}/{d.day}", count=by_day.get(d, 0)) for d in days],
        )

    by_month: dict[tuple[int, int], int] = {}
    for c in created:
        by_month[(c.year, c.month)] = by_month.get((c.year, c.month), 0) + 1
    points = []
    y, m = start.year, start.month
    while (y, m) <= (end.year, end.month):
        label = f"{m}月" if start.year == end.year else f"{y}/{m}"
        points.append(DashboardSeriesPoint(label=label, count=by_month.get((y, m), 0)))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return DashboardSeries(kind="month", label="月別", start=start, end=end, points=points)


def get_attention_count(db: Session) -> int:
    active = Reservation.status != "CANCELLED"
    return int(
        db.execute(
            select(func.count())
            .select_from(Reservation)
            .where(
                active,
                (Reservation.status == "PENDING")
                | (Reservation.has_defect.is_(True))
                | is_holiday_slot(),
            )
        ).scalar()
        or 0
    )


def build(
    db: Session,
    now: datetime | None = None,
    hospital_id: int | None = None,
    year: int | None = None,
    month: int | None = None,
    date_str: str | None = None,
    week_start: str | None = None,
    week_end: str | None = None,
) -> DashboardData:
    now = now or datetime.now()
    today = now.date()
    yesterday = today - timedelta(days=1)

    def count(*conditions) -> int:
        return int(
            db.execute(
                select(func.count()).select_from(Reservation).where(*conditions)
            ).scalar()
            or 0
        )

    active = Reservation.status != "CANCELLED"

    today_total = count(Reservation.slot_date == today, active)

    # 임시 접수와 불비, 휴진일 예약은 함께 겹칠 수 있다. 두 조건을 따로 세서 더하면
    # 같은 예약을 두 번 센다. OR 로 한 번에 센다.
    attention_count = get_attention_count(db)

    yesterday_web = count(
        func.date(Reservation.created_at) == yesterday,
        Reservation.channel == "WEB",
    )
    today_new = count(func.date(Reservation.created_at) == today)

    # --- 진행 중인 메일 발송 실패 (최신 로그가 FAILED인 건) ---
    latest_mail_subq = (
        select(
            MailLog.reservation_no,
            MailLog.template_key,
            func.max(MailLog.sent_at).label("max_sent_at"),
        )
        .where(MailLog.reservation_no != "")
        .group_by(MailLog.reservation_no, MailLog.template_key)
        .subquery()
    )

    pending_failures_stmt = (
        select(MailLog)
        .join(
            latest_mail_subq,
            (MailLog.reservation_no == latest_mail_subq.c.reservation_no)
            & (MailLog.template_key == latest_mail_subq.c.template_key)
            & (MailLog.sent_at == latest_mail_subq.c.max_sent_at)
        )
        .where(MailLog.status == "FAILED")
    )

    mail_failed = int(
        db.execute(
            select(func.count())
            .select_from(pending_failures_stmt.where(MailLog.sent_at >= now - timedelta(days=7)).subquery())
        ).scalar()
        or 0
    )

    counters = [
        DashboardCounter(
            key="attention",
            label="対応が必要な予約",
            value=attention_count,
            tone="danger" if attention_count else "neutral",
            hint="仮受付・項目漏れ・休診",
            link="#/reservations?defect=1",
        ),
        DashboardCounter(
            key="today",
            label="本日の健診",
            value=today_total,
            tone="primary",
            hint=_short_date(today),
            link=f"#/reservations?from={today.isoformat()}&to={today.isoformat()}",
        ),
        DashboardCounter(
            key="today_new",
            label="本日受付",
            value=today_new,
            hint=_short_date(today),
            link=f"#/reservations?created_date={today.isoformat()}",
        ),
        DashboardCounter(
            key="yesterday_web",
            label="昨日WEB予約",
            value=yesterday_web,
            hint=_short_date(yesterday),
            link=f"#/reservations?created_date={yesterday.isoformat()}&channel=WEB",
        ),
        DashboardCounter(
            key="mail_failed",
            label="メール送信失敗",
            value=mail_failed,
            tone="warn" if mail_failed else "neutral",
            unit="件",
            hint="直近7日",
            # 카드를 눌렀으면 실패한 건을 보러 온 것이다. 메일 문구 편집
            # 화면이 아니라 발송 이력의 「발송 실패」 목록으로 바로 보낸다.
            link="#/mail-templates?key=logs&status=FAILED",
        ),
    ]

    # --- 오늘 검진 회장별 현황 -------------------------------------------
    # 오늘 열리는 회차는 많아야 몇 건이다. 그리드 16칸을 더하는 SQL 을
    # 쓰는 대신 회차를 읽어 파이썬에서 합친다.
    today_schedules = db.execute(
        select(HospitalSchedule, Hospital)
        .join(Hospital, Hospital.id == HospitalSchedule.hospital_id)
        .where(HospitalSchedule.event_date == today)
    ).all()

    today_rows_raw = sorted(
        (
            (hospital.id, hospital.name, totals.capacity, totals.reserved)
            for schedule, hospital in today_schedules
            for totals in (grid.totals_of(schedule),)
        ),
        key=lambda row: row[3],
        reverse=True,
    )[:20]

    today_rows = [
        DashboardTodayRow(
            hospital_id=hid,
            hospital_name=name,
            reserved=int(reserved or 0),
            capacity=int(capacity or 0),
            remaining=max(0, int(capacity or 0) - int(reserved or 0)),
        )
        for hid, name, capacity, reserved in today_rows_raw
    ]

    # --- 휴진으로 닫힌 날에 남은 예약 -------------------------------------
    #
    # 마감해도 예약은 지워지지 않는다. 연락하지 않으면 그 사람들은 그 날
    # 그대로 회장에 온다. 그래서 「지금 처리할 일」에 얹는다.
    #
    # 「메일 발송 실패」만으로는 못 센다 — 이메일이 없는 사람은 발송을
    # 시도조차 하지 않아 실패로 잡히지 않는다. 세어야 할 것은 「아직
    # 연락하지 못한 사람」 전체다.
    holiday_condition = (Reservation.status != "CANCELLED") & is_holiday_slot()

    holiday = db.execute(
        select(Reservation)
        .where(holiday_condition)
        .order_by(Reservation.slot_date, Reservation.start_time)
        .limit(10)
    ).scalars().unique().all()

    holiday_total = db.execute(
        select(func.count()).select_from(Reservation).where(holiday_condition)
    ).scalar() or 0

    # 검진일이 코앞인 건. 총량보다 이쪽이 급하다 — 그 날이 오면 헛걸음한다.
    # 「대응이 필요한 예약」과 같은 기준(3일)을 쓴다.
    holiday_urgent_total = db.execute(
        select(func.count()).select_from(Reservation).where(
            holiday_condition,
            Reservation.slot_date >= today,
            Reservation.slot_date <= today + timedelta(days=3),
        )
    ).scalar() or 0

    # --- 대응이 필요한 예약 ----------------------------------------------
    attention = db.execute(
        select(Reservation)
        .where(
            Reservation.status != "CANCELLED",
            (Reservation.status == "PENDING") | (Reservation.has_defect.is_(True)),
        )
        .order_by(Reservation.slot_date, Reservation.start_time)
        .limit(10)
    ).scalars().unique().all()

    # --- 최근 접수 --------------------------------------------------------
    recent = db.execute(
        select(Reservation).order_by(Reservation.created_at.desc()).limit(10)
    ).scalars().unique().all()

    # --- 「지금 처리할 일」 집계 -------------------------------------------
    #
    # 목록은 앞의 몇 건만 내려주므로, 화면이 「30건 중 5건」이라고 적으려면
    # 전체 건수를 따로 세야 한다. 그리고 30건이 밀려 있을 때 정작 중요한 것은
    # 총량이 아니라 **검진일이 코앞인 건**이다. 검진일이 지나 버리면
    # 임시 접수인 채로 당일을 맞아 이용자가 헛걸음한다.
    attention_condition = (
        (Reservation.status != "CANCELLED")
        & ((Reservation.status == "PENDING") | (Reservation.has_defect.is_(True)))
    )
    attention_total = db.execute(
        select(func.count()).select_from(Reservation).where(attention_condition)
    ).scalar() or 0

    URGENT_DAYS = 3
    urgent_total = db.execute(
        select(func.count()).select_from(Reservation).where(
            attention_condition,
            Reservation.slot_date >= today,
            Reservation.slot_date <= today + timedelta(days=URGENT_DAYS),
        )
    ).scalar() or 0

    mail_failed_total = db.execute(
        select(func.count()).select_from(pending_failures_stmt.subquery())
    ).scalar() or 0

    # 오늘 검진이 없을 때 「그럼 언제인가」에 답한다.
    next_exam_date = db.execute(
        select(func.min(HospitalSchedule.event_date))
        .where(HospitalSchedule.event_date > today)
    ).scalar()

    # --- 메일 실패 --------------------------------------------------------
    failures = db.execute(
        pending_failures_stmt
        .order_by(MailLog.sent_at.desc())
        .limit(10)
    ).scalars().all()

    # --- 📈 실제 DB 검진 통계 집계 (성별 · 연령대별 · 접수경로 · 옵션) --------
    res_stmt = select(Reservation)
    if hospital_id:
        res_stmt = res_stmt.where(Reservation.hospital_id == hospital_id)
    res_stmt = _period_filter(
        res_stmt,
        year=year,
        month=month,
        date_str=date_str,
        week_start=week_start,
        week_end=week_end,
    )

    active_reservations = db.execute(res_stmt).scalars().all()
    total_active = len(active_reservations)

    male_cnt = sum(1 for r in active_reservations if r.gender == "M")
    female_cnt = sum(1 for r in active_reservations if r.gender == "F")

    web_cnt = sum(1 for r in active_reservations if r.channel == "WEB")
    postal_cnt = sum(1 for r in active_reservations if r.channel == "POSTAL")

    def _by(channel: str, status: str) -> int:
        return sum(1 for r in active_reservations if r.channel == channel and r.status == status)

    # 검진 대상은 만 40~74세(BR-02)다. 20·30대는 본인 확인 단계에서 걸러지므로
    # 구간을 두지 않는다. 늘 0인 행을 남겨 두면 화면과 엑셀에서 자리만 차지한다.
    age_labels = ["40代", "50代", "60代", "70代以上"]
    age_map = {lbl: {"male": 0, "female": 0} for lbl in age_labels}

    for r in active_reservations:
        age = (today - r.birth_date).days // 365 if r.birth_date else 40
        if age < 40:
            # 정상 경로로는 들어올 수 없다. 명부 오류 등으로 섞여 들어오더라도
            # 연령대 통계에는 넣지 않는다. (합계는 별도 지표를 본다)
            continue
        if age <= 49:
            lbl = "40代"
        elif age <= 59:
            lbl = "50代"
        elif age <= 69:
            lbl = "60代"
        else:
            lbl = "70代以上"

        if r.gender == "M":
            age_map[lbl]["male"] += 1
        elif r.gender == "F":
            age_map[lbl]["female"] += 1

    age_buckets = [
        DashboardAgeBucket(label=lbl, male=age_map[lbl]["male"], female=age_map[lbl]["female"])
        for lbl in age_labels
    ]

    gender_stats = DashboardGenderStats(
        male=male_cnt,
        female=female_cnt,
        age_buckets=age_buckets,
    )

    channel_stats = DashboardChannelStats(
        web=web_cnt,
        postal=postal_cnt,
        web_cancelled=_by("WEB", "CANCELLED"),
        postal_cancelled=_by("POSTAL", "CANCELLED"),
        web_pending=_by("WEB", "PENDING"),
        postal_pending=_by("POSTAL", "PENDING"),
    )

    # 상태별 (CONFIRMED / PENDING / CANCELLED) 집계도 **같은 기간**으로 센다.
    # 예전에는 기간을 무시하고 전체를 세어, 하루만 골라도 상태 막대에
    # 전체 기간의 건수가 나와 「숫자가 너무 크다」로 보였다.
    all_res_list = active_reservations

    conf_cnt = sum(1 for r in all_res_list if r.status == "CONFIRMED")
    pend_cnt = sum(1 for r in all_res_list if r.status == "PENDING")
    canc_cnt = sum(1 for r in all_res_list if r.status == "CANCELLED")

    status_stats = DashboardStatusStats(
        confirmed=conf_cnt,
        pending=pend_cnt,
        cancelled=canc_cnt,
        cancelled_advance=sum(
            1 for r in all_res_list
            if r.status == "CANCELLED" and r.cancel_type == "ADVANCE"
        ),
        cancelled_same_day=sum(
            1 for r in all_res_list
            if r.status == "CANCELLED" and r.cancel_type == "SAME_DAY"
        ),
        total=len(all_res_list),
    )

    # 옵션별 신청 건수 집계
    # 코드로만 짝을 맞춘다. 예약이 가진 이름은 신청하던 그때의 사본이라
    # 마스터에서 이름을 바꾸면 열쇠가 어긋나 건수가 통째로 0 이 된다.
    # 마스터에서 지워진 옵션의 이름을 잃지 않도록 사본도 함께 모은다.
    opt_counts = {}
    opt_names = {}
    for r in active_reservations:
        for opt in r.options:
            code = opt.option_code or "OPT"
            opt_counts[code] = opt_counts.get(code, 0) + 1
            opt_names.setdefault(code, opt.option_name)

    all_exam_options = db.execute(select(ExamOption).order_by(ExamOption.sort_order, ExamOption.id)).scalars().all()
    options_list = []

    seen_codes = set()
    for eo in all_exam_options:
        seen_codes.add(eo.code)
        cnt = opt_counts.get(eo.code, 0)
        pct = round((cnt / total_active) * 100) if total_active > 0 else 0
        options_list.append(DashboardOptionStat(code=eo.code, name=eo.name, count=cnt, percentage=pct))

    for code, cnt in opt_counts.items():
        if code not in seen_codes:
            pct = round((cnt / total_active) * 100) if total_active > 0 else 0
            options_list.append(DashboardOptionStat(
                code=code, name=opt_names.get(code, code), count=cnt, percentage=pct))

    series_start, series_end = _period_range(
        active_reservations,
        year=year,
        month=month,
        date_str=date_str,
        week_start=week_start,
        week_end=week_end,
    )
    stats_series = _stats_series(active_reservations, series_start, series_end)

    # 같은 기간의 메일 발송 실적(발송 일시 기준). 메일 기록은 예약이 자동 삭제된
    # 뒤에도 남으므로 회장으로 거르지 않고 전 회장을 센다.
    if series_start and series_end:
        mail_from = datetime(series_start.year, series_start.month, series_start.day)
        mail_to = datetime(series_end.year, series_end.month, series_end.day) + timedelta(days=1)
    else:
        mail_from, mail_to = now - timedelta(hours=24), now
    mail_period: dict[str, int] = dict(
        db.execute(
            select(MailLog.status, func.count())
            .where(MailLog.sent_at >= mail_from, MailLog.sent_at < mail_to)
            .group_by(MailLog.status)
        ).all()
    )

    # 같은 기간에 자동 삭제된 예약 수. 예약은 이미 지워졌으므로 조작 로그에
    # 남긴 건수(after_json.deleted)로 센다. (purge_service 참조)
    purge_logs = db.execute(
        select(AuditLog.after_json).where(
            AuditLog.action == "RESERVATION_PURGE",
            AuditLog.created_at >= mail_from,
            AuditLog.created_at < mail_to,
        )
    ).scalars().all()
    purge_period_deleted = sum(
        int((after or {}).get("deleted") or 0) for after in purge_logs
    )

    # 메일 실패 건의 사람·검진일. 건마다 조회하면 10번 나가므로 한 번에 받는다.
    # 예약이 이미 자동 정리로 사라진 건은 그냥 빈 값으로 둔다.
    failure_nos = [m.reservation_no for m in failures if m.reservation_no]
    failure_names: dict[str, str] = {}
    failure_dates: dict[str, date] = {}
    if failure_nos:
        for r in db.execute(
            select(Reservation).where(Reservation.reservation_no.in_(failure_nos))
        ).scalars().unique().all():
            failure_names[r.reservation_no] = r.full_name
            failure_dates[r.reservation_no] = r.slot_date

    # 최근 24시간 메일 발송 실적.
    # 화면의 「送信成功率」은 예전에 98.5% 라는 고정 문자열이었다. 메일이
    # 한 통도 나가지 않아도 98.5% 로 보고해, 운영 상태를 보는 자리에서
    # 가장 믿으면 안 되는 숫자가 되어 있었다. 실제 기록을 세어 돌려준다.
    mail_24h: dict[str, int] = dict(
        db.execute(
            select(MailLog.status, func.count())
            .where(MailLog.sent_at >= now - timedelta(hours=24))
            .group_by(MailLog.status)
        ).all()
    )

    return DashboardData(
        counters=counters,
        today=today,
        today_rows=today_rows,
        attention=[_recent(r) for r in attention],
        holiday=[_recent(r) for r in holiday],
        recent=[_recent(r) for r in recent],
        mail_failures=[
            DashboardMailFailure(
                id=m.id,
                reservation_no=m.reservation_no,
                to_email=m.to_email,
                template_label=TEMPLATE_LABELS.get(m.template_key, m.template_key),
                error_message=m.error_message,
                sent_at=m.sent_at,
                full_name=failure_names.get(m.reservation_no, ""),
                slot_date=failure_dates.get(m.reservation_no),
            )
            for m in failures
        ],
        attention_total=attention_total,
        holiday_total=holiday_total,
        holiday_urgent_total=holiday_urgent_total,
        mail_failed_total=mail_failed_total,
        urgent_total=urgent_total,
        next_exam_date=next_exam_date,
        purge_pending=purge_service.count_expired(db, now),
        purge_grace_minutes=settings.RESERVATION_EXPIRE_GRACE_MINUTES,
        mail_24h_success=mail_24h.get("SUCCESS", 0),
        mail_24h_failed=mail_24h.get("FAILED", 0),
        mail_24h_skipped=mail_24h.get("SKIPPED", 0),
        gender_stats=gender_stats,
        channel_stats=channel_stats,
        status_stats=status_stats,
        options=options_list,
        stats_series=stats_series,
        mail_period_success=mail_period.get("SUCCESS", 0),
        mail_period_failed=mail_period.get("FAILED", 0),
        mail_period_skipped=mail_period.get("SKIPPED", 0),
        purge_period_deleted=purge_period_deleted,
    )


_WEEKDAYS_JA = "月火水木金土日"


def _pct(part: int, whole: int) -> float:
    # 소수 첫째 자리까지. 정수로 반올림하면 85 + 9 + 5 = 99 처럼 합이 100 에서 어긋난다.
    return round(part * 100 / whole, 1) if whole else 0.0


# 항목별로 따로 받을 수 있는 표. 키는 파일 이름에 들어가는 영문,
# 값은 화면·파일에 적는 일본어 제목이다.
CSV_PARTS: dict[str, str] = {
    "summary": "集計",
    "age": "年代別",
    "option": "オプション",
    "time": "日別",      # 기간이 31일을 넘으면 「月別」로 바뀐다
    "venue": "会場別",
}


def export_parts_zip(
    db: Session,
    hospital_id: int | None = None,
    *,
    parts: list[str],
    year: int | None = None,
    month: int | None = None,
    date_str: str | None = None,
    week_start: str | None = None,
    week_end: str | None = None,
) -> tuple[str, bytes, str]:
    """고른 표를 각각 CSV 로 만들어 **ZIP 한 개**에 담는다.

    브라우저는 여러 파일을 한 폴더에 넣어 줄 수 없다. 한 번 내려받아 풀면
    한 폴더가 되도록 ZIP 으로 묶는다. 파일 이름은 사람이 읽는 일본어로 둔다.
    """
    period = dict(
        year=year,
        month=month,
        date_str=date_str,
        week_start=week_start,
        week_end=week_end,
    )

    buffer = io.BytesIO()
    base = ""
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for index, part in enumerate(parts, start=1):
            name, content, _ = export_file(
                db, hospital_id=hospital_id, fmt="csv", part=part, **period
            )
            base = name.rsplit("_", 1)[0]
            # ZIP 안에 폴더를 하나 두어, 풀면 파일이 흩어지지 않고 한곳에 모인다.
            archive.writestr(f"{base}/{index:02d}_{CSV_PARTS[part]}.csv", content)

    return f"{base}.zip", buffer.getvalue(), "application/zip"


def export_file(
    db: Session,
    hospital_id: int | None = None,
    fmt: str = "xlsx",
    *,
    part: str = "all",
    year: int | None = None,
    month: int | None = None,
    date_str: str | None = None,
    week_start: str | None = None,
    week_end: str | None = None,
) -> tuple[str, bytes, str]:
    """健診統計の書き出し。画面で選んだ会場・期間をそのまま使う。

    예전 CSV 는 표 네 개를 한 파일에 위아래로 이어 붙여 길고, 기간을 무시했으며,
    예약자 명세(이름·연락처)까지 들어 있었다. 표마다 탭(시트)으로 나누고
    명세는 뺀다. CSV 는 탭을 가질 수 없으므로 fmt="csv" 는 표를 두 단으로
    나란히 둔 보고서형 한 장을 낸다.
    """
    period = dict(
        year=year,
        month=month,
        date_str=date_str,
        week_start=week_start,
        week_end=week_end,
    )
    data = build(db, hospital_id=hospital_id, **period)

    hosp_name = "全会場"
    if hospital_id:
        hospital = db.get(Hospital, hospital_id)
        if hospital:
            hosp_name = hospital.name

    stmt = select(Reservation)
    if hospital_id:
        stmt = stmt.where(Reservation.hospital_id == hospital_id)
    reservations = db.execute(_period_filter(stmt, **period)).scalars().all()
    total = len(reservations)

    start, end = _period_range(reservations, **period)
    period_text = f"{start.isoformat()} ～ {end.isoformat()}" if start else "（該当なし）"
    now = datetime.now()

    g = data.gender_stats
    ch = data.channel_stats
    st = data.status_stats

    # ---- 集計 -----------------------------------------------------------
    # 취소는 한 덩어리로 세지 않는다. 事前·当日 로 나누고,
    # 이 구분이 생기기 전에 취소된 건이 있으면 그때만 한 줄 더 적어 합을 맞춘다.
    canc_unknown = st.cancelled - st.cancelled_advance - st.cancelled_same_day
    cancel_rows = [
        (CANCEL_TYPE_LABELS["ADVANCE"], st.cancelled_advance),
        (CANCEL_TYPE_LABELS["SAME_DAY"], st.cancelled_same_day),
    ]
    if canc_unknown > 0:
        cancel_rows.append(("キャンセル（種類なし・区分前）", canc_unknown))

    summary_rows = [
        ["合計", "予約件数（キャンセル含む）", total, _pct(total, total)],
        ["状態", STATUS_LABELS.get("CONFIRMED", "CONFIRMED"), st.confirmed, _pct(st.confirmed, total)],
        ["状態", STATUS_LABELS.get("PENDING", "PENDING"), st.pending, _pct(st.pending, total)],
        *[["状態", label, count, _pct(count, total)] for label, count in cancel_rows],
        # 受付経路 는 화면의 「受付の内訳」과 같게 서로 겹치지 않게 나눈다.
        # Web·郵送 은 취소를 뺀 수, 취소는 따로 — 셋을 더하면 합계.
        ["受付経路", CHANNEL_LABELS.get("WEB", "WEB"), ch.web - ch.web_cancelled,
         _pct(ch.web - ch.web_cancelled, total)],
        ["受付経路", CHANNEL_LABELS.get("POSTAL", "POSTAL"), ch.postal - ch.postal_cancelled,
         _pct(ch.postal - ch.postal_cancelled, total)],
        *[["受付経路", label, count, _pct(count, total)] for label, count in cancel_rows],
        ["性別", "男性", g.male, _pct(g.male, total)],
        ["性別", "女性", g.female, _pct(g.female, total)],
    ]


    # ---- 年代別 ---------------------------------------------------------
    age_rows = [
        [b.label, b.male, b.female, b.male + b.female, _pct(b.male + b.female, total)]
        for b in g.age_buckets
    ]

    # ---- オプション -----------------------------------------------------
    option_rows = [[o.code, o.name, o.count, _pct(o.count, total)] for o in data.options]

    # ---- 会場別 (기간 안의 예약만. 「本日」 값은 섞지 않는다) ---------------
    hospitals = db.execute(
        select(Hospital).order_by(Hospital.sort_order, Hospital.id)
    ).scalars().all()
    if hospital_id:
        hospitals = [h for h in hospitals if h.id == hospital_id]
    by_hosp: dict[int, list[Reservation]] = {}
    for r in reservations:
        by_hosp.setdefault(r.hospital_id, []).append(r)
    venue_rows = []
    for h in hospitals:
        rs = by_hosp.get(h.id, [])
        venue_rows.append([
            h.code,
            h.name,
            len(rs),
            sum(1 for r in rs if r.status == "CONFIRMED"),
            sum(1 for r in rs if r.status == "PENDING"),
            sum(1 for r in rs if r.status == "CANCELLED" and r.cancel_type == "ADVANCE"),
            sum(1 for r in rs if r.status == "CANCELLED" and r.cancel_type == "SAME_DAY"),
            sum(1 for r in rs if r.channel == "WEB" and r.status != "CANCELLED"),
            sum(1 for r in rs if r.channel == "POSTAL" and r.status != "CANCELLED"),
            sum(1 for r in rs if r.gender == "M"),
            sum(1 for r in rs if r.gender == "F"),
        ])

    # ---- 日別 (기간 안의 모든 날짜. 접수가 없는 날도 0 으로 남긴다) --------
    by_day: dict[date, list[Reservation]] = {}
    for r in reservations:
        if r.created_at:
            by_day.setdefault(r.created_at.date(), []).append(r)
    day_rows = []
    if start and end:
        day = start
        while day <= end:
            rs = by_day.get(day, [])
            day_rows.append([
                day.isoformat(),
                _WEEKDAYS_JA[day.weekday()],
                len(rs),
                sum(1 for r in rs if r.status == "CONFIRMED"),
                sum(1 for r in rs if r.status == "PENDING"),
                sum(1 for r in rs if r.status == "CANCELLED" and r.cancel_type == "ADVANCE"),
                sum(1 for r in rs if r.status == "CANCELLED" and r.cancel_type == "SAME_DAY"),
                sum(1 for r in rs if r.channel == "WEB" and r.status != "CANCELLED"),
                sum(1 for r in rs if r.channel == "POSTAL" and r.status != "CANCELLED"),
                sum(1 for r in rs if r.gender == "M"),
                sum(1 for r in rs if r.gender == "F"),
            ])
            day += timedelta(days=1)

    # ---- 日別 또는 月別 (기간이 31일을 넘으면 월로 묶는다) ----------------
    span_days = (end - start).days + 1 if start and end else 0
    if span_days > 31:
        by_month: dict[str, list[Reservation]] = {}
        for r in reservations:
            if r.created_at:
                by_month.setdefault(r.created_at.strftime("%Y-%m"), []).append(r)
        months = []
        y, m = start.year, start.month
        while (y, m) <= (end.year, end.month):
            months.append(f"{y:04d}-{m:02d}")
            y, m = (y + 1, 1) if m == 12 else (y, m + 1)
        time_title = "月別（期間が31日を超えるため月ごと）"
        time_header_head = ["年月", ""]
        time_rows = []
        for key in months:
            rs = by_month.get(key, [])
            time_rows.append([
                key, "", len(rs),
                sum(1 for r in rs if r.status == "CONFIRMED"),
                sum(1 for r in rs if r.status == "PENDING"),
                sum(1 for r in rs if r.status == "CANCELLED" and r.cancel_type == "ADVANCE"),
                sum(1 for r in rs if r.status == "CANCELLED" and r.cancel_type == "SAME_DAY"),
                sum(1 for r in rs if r.channel == "WEB" and r.status != "CANCELLED"),
                sum(1 for r in rs if r.channel == "POSTAL" and r.status != "CANCELLED"),
                sum(1 for r in rs if r.gender == "M"),
                sum(1 for r in rs if r.gender == "F"),
            ])
    else:
        time_title = "日別"
        time_header_head = ["日付", "曜日"]
        time_rows = day_rows

    stamp = f"{start:%Y%m%d}-{end:%Y%m%d}" if start else now.strftime("%Y%m%d")
    base_name = f"health_stats_{stamp}_{'all' if not hospital_id else hospital_id}"

    # 통계에서는 「キャンセル」 한 덩어리를 쓰지 않는다. 事前·当日 로 나눠 센다.
    status_head = [
        STATUS_LABELS.get("CONFIRMED", "CONFIRMED"),
        STATUS_LABELS.get("PENDING", "PENDING"),
        CANCEL_TYPE_LABELS["ADVANCE"],
        CANCEL_TYPE_LABELS["SAME_DAY"],
        CHANNEL_LABELS.get("WEB", "WEB"),
        CHANNEL_LABELS.get("POSTAL", "POSTAL"),
        "男性",
        "女性",
    ]

    # ---- 항목별 CSV 한 장 -------------------------------------------------
    # 「필요한 표만 골라 받고 싶다」는 요청에 맞춘다. 한 파일에 표 하나만 담아
    # 다른 곳(보고서·집계표)에 그대로 붙여 쓸 수 있게 한다.
    if part != "all":
        if part not in CSV_PARTS:
            raise ValueError(f"unknown part: {part}")

        if part == "summary":
            title = "集計"
            header = ["区分", "項目", "件数", "割合(%)"]
            body_rows = [[k, item, n, p] for k, item, n, p in summary_rows]
        elif part == "age":
            title = "年代別"
            header = ["年代", "男性", "女性", "合計", "割合(%)"]
            body_rows = age_rows
        elif part == "option":
            title = "オプション"
            header = ["コード", "検査名", "件数", "割合(%)"]
            body_rows = option_rows
        elif part == "time":
            title = time_title
            header = time_header_head + ["件数"] + status_head
            body_rows = time_rows
        else:
            title = "会場別"
            header = ["会場コード", "会場名", "件数"] + status_head
            body_rows = venue_rows

        output = io.StringIO()
        writer = csv.writer(output, lineterminator="\r\n")
        writer.writerow(["【期間】", period_text, "【会場】", hosp_name,
                         "【基準】", "申込日（キャンセル含む）",
                         "【対象件数】", total, "【出力】", now.strftime("%Y-%m-%d %H:%M")])
        writer.writerow([])
        writer.writerow(["【" + title + "】"])
        writer.writerow([f"【{h}】" if h else "" for h in header])
        writer.writerows(body_rows)
        return (
            f"{base_name}_{part}.csv",
            output.getvalue().encode("utf-8-sig"),
            "text/csv; charset=utf-8-sig",
        )

    if fmt == "xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            fmt = "csv"
        else:
            wb = Workbook()
            head_font = Font(bold=True, color="FFFFFFFF")
            head_fill = PatternFill("solid", fgColor="FF0B6E5B")

            def sheet(title, header, rows, widths, first=False):
                ws = wb.active if first else wb.create_sheet()
                ws.title = title
                ws.append(header)
                for i in range(1, len(header) + 1):
                    cell = ws.cell(row=1, column=i)
                    cell.font = head_font
                    cell.fill = head_fill
                for row in rows:
                    ws.append(row)
                ws.freeze_panes = "A2"
                if rows:
                    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
                for i, w in enumerate(widths, start=1):
                    ws.column_dimensions[get_column_letter(i)].width = w
                return ws

            sheet("条件", ["項目", "内容"], [
                ["出力日時", now.strftime("%Y-%m-%d %H:%M:%S")],
                ["会場", hosp_name],
                ["集計の基準", "申込日（受付日）"],
                ["期間", period_text],
                ["対象件数", total],
            ], [18, 40], first=True)
            sheet("集計", ["区分", "項目", "件数", "割合(%)"], summary_rows, [12, 30, 10, 10])
            sheet("日別", ["日付", "曜日", "件数"] + status_head, day_rows,
                  [12, 6, 8, 10, 12, 10, 16, 10, 8, 8])
            sheet("会場別", ["会場コード", "会場名", "件数"] + status_head, venue_rows,
                  [12, 36, 8, 10, 12, 10, 16, 10, 8, 8])
            sheet("年代別", ["年代", "男性", "女性", "合計", "割合(%)"], age_rows, [12, 8, 8, 8, 10])
            sheet("オプション", ["コード", "検査名", "件数", "割合(%)"], option_rows, [10, 34, 8, 10])
            sheet("本日の運営状況", ["項目", "数値", "単位", "備考"],
                  [[c.label, c.value, c.unit, c.hint] for c in data.counters],
                  [22, 10, 8, 30])

            stream = io.BytesIO()
            wb.save(stream)
            return (
                base_name + ".xlsx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

    # ---- CSV: 보고서형 (표를 옆으로 나란히) ----------------------------------
    # CSV 는 탭도 열 너비도 없다. 표를 위아래로 이으면 100 줄이 넘어가므로
    # 두 단으로 나란히 둔다. 긴 글자(항목명·검사명·회장명)는 블록 맨 오른쪽에
    # 두고 그 오른쪽을 비워, 엑셀에서 열었을 때 글자가 넘쳐서라도 다 보이게 한다.
    # 기간이 31일을 넘으면 날짜별은 월별로 묶어 세로 길이를 억누른다.
    # 会場別は「件数＋状態8列＋会場コード・会場名」で11列を使う。
    # 取り消しを 事前／当日 に分けた分だけ横に伸びたので、幅もそれに合わせる。
    width = 24  # A..X（最後の1列は会場名がはみ出して見えるように空ける）
    grid: list[list] = []

    def put(r: int, col: int, value) -> None:
        while len(grid) <= r:
            grid.append([""] * width)
        grid[r][col] = "" if value is None else value

    def place(top: int, col: int, title: str, header: list, rows: list) -> int:
        put(top, col, "【" + title + "】")
        for i, h in enumerate(header):
            put(top + 1, col + i, h)
        for r_i, row in enumerate(rows):
            for c_i, v in enumerate(row):
                put(top + 2 + r_i, col + c_i, v)
        return top + 2 + len(rows)

    # 1행의 값도 오른쪽 칸을 비워 두어 잘리지 않게 한다.
    put(0, 0, "【期間】")
    put(0, 1, period_text)        # B (C·D 비움)
    put(0, 4, "【会場】")
    put(0, 5, hosp_name)          # F (G 비움)
    put(0, 7, "【基準】")
    put(0, 8, "申込日（キャンセル含む）")  # I (J·K 비움)
    put(0, 11, "【対象件数】")
    put(0, 12, total)
    put(0, 14, "【出力】")
    put(0, 15, now.strftime("%Y-%m-%d %H:%M"))  # P

    short_label = {"予約件数（キャンセル含む）": "予約件数"}
    csv_summary = [[k, n, p, short_label.get(item, item)] for k, item, n, p in summary_rows]
    csv_options = [[code, n, p, name] for code, name, n, p in option_rows]

    time_header = time_header_head + ["件数"] + status_head

    csv_venues = [row[2:] + [row[0], row[1]] for row in venue_rows]

    end1 = max(
        place(2, 0, "集計", ["区分", "件数", "割合(%)", "項目"], csv_summary),
        place(2, 5, "年代別", ["年代", "男性", "女性", "合計", "割合(%)"], age_rows),
        place(2, 11, "オプション", ["コード", "件数", "割合(%)", "検査名"], csv_options),
    )
    top2 = end1 + 1
    place(top2, 0, time_title, time_header, time_rows)
    place(top2, 12, "会場別", ["件数"] + status_head + ["会場コード", "会場名"], csv_venues)

    output = io.StringIO()
    # BOM 은 여기서 붙이지 않는다 — 아래 encode("utf-8-sig") 가 붙인다.
    csv.writer(output, lineterminator="\r\n").writerows(grid)
    return base_name + ".csv", output.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8-sig"

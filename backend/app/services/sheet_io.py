"""스프레드시트 표 ↔ CSV / Excel 변환 엔진.

관리 화면의 「일괄 입력」 표는 지금 두 곳에 있다.

  · 우편 접수 일괄 입력 (A-13) — `postal_sheet_service`
  · 회장 일괄 관리   (A-24) — `hospital_sheet_service`

둘이 다른 것은 **열의 정의뿐**이다. 제목 줄을 찾는 법, 일본어 Excel 이
저장한 CP932 를 알아보는 법, 우편번호의 앞자리 0 이 날아간 것을 눈치채는
법은 완전히 같다. 그 공통부를 여기 한 곳에 둔다.

같은 코드를 두 벌 두면 한쪽에서 고친 버그가 다른 쪽에 남는다. 특히 이
모듈이 다루는 것들 — 인코딩 판정, 수식 인젝션 방어, Excel 의 날짜 셀 —
은 전부 「한 번 제대로 짜고 다시 건드리지 않는」 종류의 코드다.

쓰는 쪽은 `SheetFormat` 을 하나 만들어 두고 그 메서드를 부른다.

    FORMAT = SheetFormat(
        columns=(SheetColumn("code", "코드", 12), ...),
        sheet_name="회장",
        file_prefix="회장_일괄관리",
    )
    FORMAT.build_xlsx(rows)
    FORMAT.parse_upload(name, raw)

업무 규칙은 여기서 검증하지 않는다
----------------------------------
형식이 틀린 셀도 그대로 표에 올려 사람이 보고 고치게 한다. 판정은 각
화면의 일괄 검증 한 곳에서만 한다. 여기서 한 번 더 거절하면 같은 값이
화면에서는 통과하고 파일에서는 막히는 어긋남이 생긴다.

다만 **형식 변환은 한다.** Excel 의 날짜 셀은 문자열이 아니라 `datetime`
으로 넘어오므로, 표가 쓰는 표기(`YYYY-MM-DD`)로 바꿔 주지 않으면 사람이
파일에서 본 것과 표에 뜬 것이 달라진다.
"""

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any

# 표(프런트의 `MAX_ROWS`)와 같은 상한. 어느 한쪽만 늘리면
# 「파일은 읽혔는데 표에 안 들어간다」가 된다.
MAX_SHEET_ROWS = 200

# 업로드 상한. 200행 CSV 는 100KB 수준이라 사실상 오타·오파일 방어선이다.
MAX_UPLOAD_BYTES = 2 * 1024 * 1024

# 한 셀의 길이 상한. 메모(10,000자)가 가장 긴 항목이다.
MAX_CELL_LENGTH = 10_000

# 제목 줄을 찾기 위해 훑는 최대 줄 수. 표 위에 「2026년도 접수분」 같은
# 제목 한 줄이 붙어 있는 서식이 흔하다.
HEADER_SCAN_ROWS = 8

# 파일에 담는 참/거짓 표기. 표의 표시값과 같아야 왕복이 성립한다.
BOOL_TRUE_TEXT = "はい"
BOOL_FALSE_TEXT = "いいえ"

GUIDE_SHEET_NAME = "入力案内"

# 한국어 시절의 열 이름. 화면 문구를 일본어로 바꾸면서 열 머리글도 함께
# 바뀌었는데, **담당자 손에 있는 파일은 그대로다.** 예전에 내보낸 표를 다시
# 올리면 열이 하나도 붙지 않아 「형식이 다릅니다」로 끝난다.
#
# 그래서 옛 이름을 별칭으로 계속 받는다. 열 정의(SheetColumn)에 섞어 두지
# 않고 여기 한 곳에 모으는 이유는, 이것이 번역의 일부가 아니라 **한 번
# 지나가면 지워도 되는 하위 호환**이기 때문이다. 어느 날 지울 때 이 표만
# 지우면 된다.
LEGACY_HEADERS: dict[str, tuple[str, ...]] = {
    "code": ("회장 코드", "코드", "검사코드",),
    "name": ("회장명", "검사명", "옵션명",),
    "area": ("지역", "관할",),
    "city": ("시구정촌", "시군구",),
    "postal_code": ("우편번호",),
    "address": ("주소",),
    "transit_info": ("교통정보", "최기역",),
    "access_minutes": ("소요시간", "도보",),
    "event_date": ("개최일", "검진일",),
    "booking_close_date": ("접수 마감일", "마감일",),
    "open_time": ("개시 시각", "개시",),
    "reception_end_time": ("접수 종료 시각", "접수 종료",),
    "schedule_note": ("회차 메모", "메모", "비고",),
    "has_parking": ("주차장",),
    "latitude": ("위도",),
    "longitude": ("경도",),
    "tel": ("전화(예외)", "전화", "전화번호",),
    "is_visible": ("예약 화면 표시", "표시", "공개", "표시여부",),
    "sort_order": ("정렬 순서", "순서", "정렬", "정렬순서",),
    "hospital_code": ("회장 코드", "코드",),
    "hospital_name": ("회장명",),
    "note": ("회차 메모", "메모", "비고", "준비 사항", "준비사항", "주의 사항",),
    "last_name": ("성", "성(한자)",),
    "first_name": ("이름", "이름(한자)",),
    "last_name_kana": ("성 후리가나", "성카나",),
    "first_name_kana": ("이름 후리가나", "이름카나",),
    "middle_name": ("미들네임",),
    "middle_name_kana": ("미들네임 후리가나",),
    "gender": ("성별",),
    "birth_date": ("생년월일",),
    "insurer_no": ("보험자 번호",),
    "insurance_symbol": ("보험증 기호",),
    "insurance_no": ("보험증 번호",),
    "address_detail": ("번지",),
    "building": ("아파트·맨션명", "건물명",),
    "tel_mobile": ("휴대전화",),
    "tel_home": ("유선전화",),
    "email": ("이메일", "메일",),
    "wish1_hospital": ("제1희망 회장", "회장1",),
    "wish1_date": ("제1희망 검진일", "검진일1",),
    "wish1_time": ("제1희망 시작시각", "시각1",),
    "wish2_hospital": ("제2희망 회장", "회장2",),
    "wish2_date": ("제2희망 검진일", "검진일2",),
    "wish2_time": ("제2희망 시작시각", "시각2",),
    "wish3_hospital": ("제3희망 회장", "회장3",),
    "wish3_date": ("제3희망 검진일", "검진일3",),
    "wish3_time": ("제3희망 시작시각", "시각3",),
    "option_codes": ("옵션 코드", "옵션",),
    "allow_defect": ("불비 허용", "불비",),
    "memo": ("스태프 메모", "메모", "비고",),
    "is_active": ("사용 여부", "사용여부", "활성", "사용함",),
    "target_gender": ("타깃 성별", "성별",),
    "target_age_min": ("최소 연령", "최소연령",),
    "target_age_max": ("최대 연령", "최대연령",),
    "description": ("설명", "안내",),
    "name_kana": ("후리가나",),
}



class SheetError(Exception):
    """파일을 표로 바꿀 수 없는 이유. API 는 400 으로 돌려준다."""


@dataclass(frozen=True)
class SheetColumn:
    key: str
    label: str
    width: int
    kind: str = "text"          # text / date / time / bool / number
    aliases: tuple[str, ...] = ()
    # 앞자리 0 이 날아가면 값이 달라지는 칸 (우편번호·전화번호 등).
    # 숫자 셀로 들어오면 경고한다.
    zero_sensitive: bool = False
    # 한 셀에 값을 여러 개 적는 칸. 줄바꿈을 세미콜론으로 살린다.
    multi: bool = False


@dataclass
class SheetImport:
    rows: list[dict[str, str]] = field(default_factory=list)
    kind: str = "CSV"
    encoding: str = ""
    sheet_name: str = ""
    header_mode: str = "HEADER"
    blank_rows: int = 0
    matched_columns: int = 0
    ignored_columns: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def normalize_header(value: Any) -> str:
    """제목 줄 대조용 정규화.

    전각·공백·괄호·구분 기호를 지우고 접어 비교한다. 이 덕에
    `제1희망 회장` · `제 1 희망 회장` · `第１希望病院` · `wish1_hospital`
    이 모두 같은 열로 붙는다.
    """
    text = unicodedata.normalize("NFKC", "" if value is None else str(value))
    text = text.replace("\ufeff", "")  # BOM 제거
    text = re.sub(r"[\s()\[\]{}【】·・:：/／\\\-_.,'\"]", "", text)
    return text.casefold()


def _openpyxl():
    """Excel 을 다루는 순간에만 불러온다.

    설치되지 않은 환경에서 이 모듈을 import 하는 것만으로 관리 API 전체가
    죽는 것을 막는다. CSV 는 표준 라이브러리만으로 되므로 그때까지는
    openpyxl 이 없어도 화면이 동작해야 한다.
    """
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:  # pragma: no cover - 설치 누락 환경
        raise SheetError(
            "Excel ファイルを扱うにはサーバーに openpyxl が必要です。"
            "担当者にお問い合わせください。(pip install -r requirements.txt)"
        ) from exc
    return Workbook, load_workbook


class SheetFormat:
    """열 정의 한 벌 + 그 열로 하는 파일 입출력."""

    def __init__(
        self,
        *,
        columns: tuple[SheetColumn, ...],
        sheet_name: str,
        file_prefix: str,
        guide_lines: tuple[tuple[str, str], ...] = (),
        max_rows: int = MAX_SHEET_ROWS,
        dynamic_column_pattern: re.Pattern | None = None,
    ) -> None:
        self.columns = columns
        self.sheet_name = sheet_name
        self.file_prefix = file_prefix
        self.guide_lines = guide_lines
        self.max_rows = max_rows
        self.dynamic_column_pattern = dynamic_column_pattern

        self.keys: tuple[str, ...] = tuple(column.key for column in columns)
        self._by_key: dict[str, SheetColumn] = {c.key: c for c in columns}

        # 열 키·라벨·별칭 어느 것으로도 붙게 한다. 먼저 등록된 쪽이 이긴다.
        self.alias_to_key: dict[str, str] = {}
        # 같은 열 키에 붙는 이름들 사이의 순위. 작을수록 그 열을 **정확히** 뜻한다.
        # 한 제목 줄에 둘 이상이 함께 있을 때 어느 칸을 쓸지 여기서 정한다.
        self.alias_rank: dict[str, int] = {}
        for column in columns:
            for rank, candidate in enumerate((column.key, column.label, *column.aliases)):
                normalized = normalize_header(candidate)
                self.alias_to_key.setdefault(normalized, column.key)
                self.alias_rank.setdefault(normalized, rank)

        # 옛 한국어 머리글은 **맨 뒤에** 넣는다. 지금 쓰는 이름과 겹치는 일이
        # 있어도 현재 정의가 이기게 하기 위해서다.
        for column in columns:
            for candidate in LEGACY_HEADERS.get(column.key, ()):
                normalized = normalize_header(candidate)
                self.alias_to_key.setdefault(normalized, column.key)
                self.alias_rank.setdefault(normalized, 1000)

    # ----------------------------------------------------------------------
    # 셀 값 변환
    # ----------------------------------------------------------------------

    def empty_row(self) -> dict[str, str]:
        return {key: "" for key in self.keys}

    @staticmethod
    def _number_text(value: float | int) -> str:
        """숫자 셀을 사람이 적었을 모습으로 되돌린다.

        Excel 은 정수도 실수로 들고 있어 `1010021.0` 처럼 넘어온다.
        그대로 두면 우편번호가 소수점을 달고 표에 들어간다.
        """
        if isinstance(value, bool):
            return BOOL_TRUE_TEXT if value else ""
        if isinstance(value, int):
            return str(value)
        if value == int(value):
            return str(int(value))
        return repr(value)

    @staticmethod
    def _fraction_to_time(value: float) -> str:
        """0 이상 1 미만의 실수를 시각으로 본다. Excel 의 시각 표현이다."""
        minutes = round(value * 24 * 60)
        return f"{(minutes // 60) % 24:02d}:{minutes % 60:02d}"

    def cell_to_text(self, value: Any, column: SheetColumn) -> tuple[str, bool]:
        """셀 하나를 표가 쓰는 문자열로 바꾼다.

        돌려주는 두 번째 값은 「숫자 셀이었는가」다. 전화번호·우편번호가
        숫자 셀로 들어오면 앞자리 0 이 이미 사라진 상태이므로 경고해야 한다.
        """
        if value is None:
            return "", False

        if isinstance(value, bool):
            return (BOOL_TRUE_TEXT if value else ""), False

        if isinstance(value, datetime):
            if column.kind == "time":
                return value.strftime("%H:%M"), False
            return value.strftime("%Y-%m-%d"), False

        if isinstance(value, date):
            return value.strftime("%Y-%m-%d"), False

        if isinstance(value, time):
            return value.strftime("%H:%M"), False

        if isinstance(value, (int, float)):
            if column.kind == "time" and 0 <= float(value) < 1:
                return self._fraction_to_time(float(value)), False
            return self._number_text(value), True

        text = str(value)

        # 내보낼 때 붙인 수식 방어용 따옴표를 걷어 낸다. 붙인 쪽과 떼는 쪽이
        # 같은 규칙을 쓰지 않으면 왕복할 때마다 따옴표가 하나씩 쌓인다.
        if len(text) > 1 and text[0] == "'" and text[1] in "=+-@":
            text = text[1:]

        # 표의 셀은 한 줄이다. 줄바꿈을 그대로 넣으면 표시가 깨진다.
        # 값을 여러 개 적는 칸만은 줄바꿈을 구분자로 살려 준다 — 한 셀에
        # 코드를 줄마다 적어 두는 서식이 흔하고, 공백으로 이으면 한 덩어리가
        # 되어 버린다.
        if column.multi:
            text = re.sub(r"[\r\n]+", ";", text)
        else:
            text = re.sub(r"[\r\n\t]+", " ", text)

        text = text.strip()
        if len(text) > MAX_CELL_LENGTH:
            text = text[:MAX_CELL_LENGTH]
        return text, False

    @staticmethod
    def export_text(value: Any) -> str:
        """표에서 받은 값을 파일에 적을 문자열로 바꾼다.

        붙여넣기로 흔히 섞여 오는 줄바꿈 없는 공백(NBSP)은 보통 공백으로 맞춘다.
        눈에는 같은 공백이라 「왜 검색이 안 되는가」로만 드러난다.
        """
        if value is None or value is False:
            return ""
        if value is True:
            return BOOL_TRUE_TEXT
        text = str(value).replace(" ", " ")
        text = re.sub(r"[\r\n\t]+", " ", text).strip()
        if len(text) > MAX_CELL_LENGTH:
            text = text[:MAX_CELL_LENGTH]
        return text

    @staticmethod
    def _guard_formula(text: str) -> str:
        """`=`·`+`·`-`·`@` 로 시작하는 셀을 수식으로 실행하지 않게 한다.

        내려받은 파일을 Excel 에서 여는 것이 이 기능의 목적이므로, 셀 값이
        그대로 수식이 되는 경로(CSV 인젝션)를 막아 둔다. 앞에 붙인 따옴표는
        다시 가져올 때 `cell_to_text` 가 떼어 낸다.
        """
        return "'" + text if text[:1] in ("=", "+", "-", "@") else text

    def normalize_export_rows(
        self, rows: list[dict[str, Any]]
    ) -> list[dict[str, str]]:
        """화면에서 받은 행을 열 키만 남긴 문자열 행으로 정리한다.

        모르는 열은 조용히 버린다. 화면이 열을 하나 늘렸는데 이 모듈이 아직
        모른다면, 그것은 내보내기가 실패할 일이 아니라 다음 배포에서 열을
        추가할 일이다.
        """
        cleaned: list[dict[str, str]] = []
        for raw in rows:
            row = {key: self.export_text(raw.get(key)) for key in self.keys}
            if any(row.values()):
                cleaned.append(row)
        return cleaned

    # ----------------------------------------------------------------------
    # 내보내기
    # ----------------------------------------------------------------------

    def build_csv(self, rows: list[dict[str, str]]) -> bytes:
        """UTF-8 BOM + CRLF.

        BOM 이 없으면 일본어·한국어 Windows 의 Excel 이 CSV 를 ANSI 로 열어
        제목 줄부터 깨진 채 뜬다. 「Excel 에서 바로 열린다」가 이 기능의
        전부이므로 BOM 을 붙인다.
        """
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer, lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
        # CSV 는 색·굵기가 없으므로 머리글을 【】로 감싼다. 가져올 때는
        # normalize_header 가 【】를 떼므로 괄호 없는 예전 파일도 그대로 붙는다.
        writer.writerow([f"【{column.label}】" for column in self.columns])
        for row in rows:
            writer.writerow(
                [self._guard_formula(row.get(c.key, "")) for c in self.columns]
            )
        return buffer.getvalue().encode("utf-8-sig")

    def build_xlsx(self, rows: list[dict[str, str]]) -> bytes:
        """Excel 파일 1장(데이터) + 1장(입력 안내).

        모든 셀을 **문자열 서식**으로 쓴다. 그렇지 않으면 Excel 이 우편번호
        `0110021` 의 앞자리 0 을 지우고, `2026-08-12` 를 자기 지역 표기의
        날짜로 바꿔 버린다. 그 파일을 다시 가져오면 값이 달라져 있다.
        """
        Workbook, _load_workbook = _openpyxl()
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        book = Workbook()
        sheet = book.active
        sheet.title = self.sheet_name

        header_font = Font(bold=True, color="FF1F3B34")
        header_fill = PatternFill("solid", fgColor="FFE6EFED")
        header_align = Alignment(vertical="center", wrap_text=False)

        sheet.append([column.label for column in self.columns])
        for index, column in enumerate(self.columns, start=1):
            cell = sheet.cell(row=1, column=index)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            sheet.column_dimensions[get_column_letter(index)].width = column.width

        for row_index, row in enumerate(rows, start=2):
            for column_index, column in enumerate(self.columns, start=1):
                cell = sheet.cell(row=row_index, column=column_index)
                text = row.get(column.key, "")
                cell.value = text
                # openpyxl 은 `=` 로 시작하는 문자열을 **수식**으로 저장한다.
                # 그러면 계산 결과가 없는 셀이 되어 다시 가져올 때 값이
                # 빈칸으로 사라진다. 문자열 셀임을 못 박아 원문을 지킨다.
                if text.startswith("="):
                    cell.data_type = "s"
                cell.number_format = "@"

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(self.columns))}{max(1, len(rows) + 1)}"
        )

        if self.guide_lines:
            guide = book.create_sheet(GUIDE_SHEET_NAME)
            guide.append(["項目", "入力方法"])
            for label, description in self.guide_lines:
                guide.append([label, description])
            guide.column_dimensions["A"].width = 22
            guide.column_dimensions["B"].width = 76
            for index in range(1, len(self.guide_lines) + 2):
                guide.cell(row=index, column=1).font = Font(bold=index == 1)
                guide.cell(row=index, column=2).alignment = Alignment(wrap_text=True)

        stream = io.BytesIO()
        book.save(stream)
        return stream.getvalue()

    def export_file_name(self, kind: str, moment: datetime) -> str:
        """내려받을 파일 이름. 같은 날 여러 번 받아도 섞이지 않게 분까지 넣는다."""
        suffix = "xlsx" if kind == "xlsx" else "csv"
        return f"{self.file_prefix}_{moment.strftime('%Y%m%d_%H%M')}.{suffix}"

    # ----------------------------------------------------------------------
    # 가져오기
    # ----------------------------------------------------------------------

    def _header_hits(self, cells: list[Any]) -> tuple[int, dict[int, str]]:
        """이 줄이 제목 줄이라면 어느 칸이 어느 열인지 짝지어 본다."""
        mapping: dict[int, str] = {}
        # \uc5f4 \ud0a4 \u2192 (\uadf8 \ud0a4\ub97c \ucc28\uc9c0\ud55c \uce78, \uc21c\uc704)
        taken: dict[str, tuple[int, int]] = {}
        for index, cell in enumerate(cells):
            text = str(cell or "").strip()
            text = text.replace("\ufeff", "")
            normalized = normalize_header(text)
            key = self.alias_to_key.get(normalized)
            if key:
                # \uac19\uc740 \uc5f4\ub85c \uc77d\ud788\ub294 \uce78\uc774 \ub458\uc774\uba74 **\uba3c\uc800 \ub098\uc628 \uce78\uc774 \uc544\ub2c8\ub77c \ub354 \uc815\ud655\ud55c
                # \uc774\ub984**\uc744 \uc4f4\ub2e4. \uc6d0\ubcf8 \ub9c8\uc2a4\ud130\ub294 `ID`(\ud589 \ubc88\ud638)\uac00 `\u4f1a\u5834\u756a\u53f7` \ubcf4\ub2e4
                # \uc55e\uc5d0 \uc788\ub294\ub370, \uba3c\uc800 \ub098\uc628 \ucabd\uc744 \uc4f0\uba74 \ud68c\uc7a5 \ucf54\ub4dc T01 \uc774 \ud589 \ubc88\ud638 1 \u2192
                # V01 \ub85c \uc77d\ud600 **\uc5c9\ub6b1\ud55c \ud68c\uc7a5\uc744 \ub36e\uc5b4\uc4f4\ub2e4.**
                rank = self.alias_rank.get(normalized, 999)
                previous = taken.get(key)
                if previous is None or rank < previous[1]:
                    if previous is not None:
                        mapping.pop(previous[0], None)
                    mapping[index] = key
                    taken[key] = (index, rank)
            elif self.dynamic_column_pattern and self.dynamic_column_pattern.match(text):
                # 동적 패턴과 일치하는 열(예: 정원 슬롯)은 원본 값을 그대로 키로 사용한다.
                mapping[index] = text
        return len(mapping), mapping

    def _find_header(self, matrix: list[list[Any]]) -> tuple[int, dict[int, str]]:
        """제목 줄을 찾는다. 못 찾으면 `(-1, {})`.

        표 위에 제목 한 줄이 붙어 있는 서식이 흔하므로 앞쪽 몇 줄을 훑는다.
        2칸 이상 맞으면 제목 줄로 본다. 1칸만 맞는 줄을 제목으로 보면
        「이름」이 적힌 데이터 행을 제목으로 오인해 그 행을 잃는다.
        """
        best_index = -1
        best_mapping: dict[int, str] = {}
        for index, cells in enumerate(matrix[:HEADER_SCAN_ROWS]):
            hits, mapping = self._header_hits(cells)
            if hits >= 2 and hits > len(best_mapping):
                best_index = index
                best_mapping = mapping
        return best_index, best_mapping

    def _is_blank(self, cells: list[Any]) -> bool:
        return not any(self.export_text(cell) for cell in cells)

    def _build_rows(self, matrix: list[list[Any]], result: SheetImport) -> None:
        """행렬을 열 키 기반 행으로 옮긴다. 제목 줄 인식은 여기서 한다."""
        header_index, mapping = self._find_header(matrix)

        if header_index >= 0:
            result.header_mode = "HEADER"
            header_cells = matrix[header_index]
            body = matrix[header_index + 1 :]
            result.ignored_columns = [
                self.export_text(cell)
                for index, cell in enumerate(header_cells)
                if index not in mapping and self.export_text(cell)
            ]
        else:
            # 제목 줄이 없어도, 열 수가 정확히 맞으면 서식 그대로라고 보고 받는다.
            # 열 수까지 다르면 추측하지 않는다. 잘못 맞추면 성과 이름이,
            # 우편번호와 전화번호가 뒤바뀐 값이 조용히 저장된다.
            width = max((len(cells) for cells in matrix), default=0)
            if width != len(self.columns):
                raise SheetError(
                    "ヘッダー行が見つかりません。「Excel書き出し」や"
                    "「CSV 書き出し」で取得したフォーマットのヘッダー行をそのまま1行目に"
                    f"設定してください。（ヘッダー行がない場合は列が正確に "
                    f"{len(self.columns)}列である必要があります）"
                )
            result.header_mode = "POSITIONAL"
            mapping = {index: key for index, key in enumerate(self.keys)}
            body = matrix

        result.matched_columns = len(mapping)

        numeric_hits = 0
        rows: list[dict[str, str]] = []
        for cells in body:
            if self._is_blank(cells):
                result.blank_rows += 1
                continue
            if len(rows) >= self.max_rows:
                raise SheetError(
                    f"データ行が {self.max_rows}行を超えています。"
                    f"一括で取り込み可能な行数は {self.max_rows}行までのため、"
                    "ファイルを分割してアップロードしてください。"
                )

            row = self.empty_row()
            for index, key in mapping.items():
                if index >= len(cells):
                    continue
                if key in self._by_key:
                    column = self._by_key[key]
                    text, was_numeric = self.cell_to_text(cells[index], column)
                    row[key] = text
                    if was_numeric and column.zero_sensitive and text:
                        numeric_hits += 1
                else:
                    # 동적 컬럼인 경우 문자열로만 변환하여 추가한다.
                    row[key] = self.export_text(cells[index])
            rows.append(row)

        result.rows = rows

        if numeric_hits:
            result.warnings.append(
                f"郵便番号・電話番号のように先頭の「0」が重要な項目 {numeric_hits}箇所が"
                "数値セルとして保存されていました。先頭の「0」が欠落している可能性があるため、"
                "一覧でご確認ください。（Excelで該当列を「文字列」書式に"
                "設定しておくと発生しません）"
            )
        if not rows:
            raise SheetError(
                "取り込むデータ行がありません。ヘッダー行の下に内容があるか"
                "ご確認ください。"
            )

    def _decode_csv(self, raw: bytes) -> tuple[str, str]:
        """CSV 의 인코딩을 정한다.

        일본어 Excel 의 「CSV(쉼표 구분)」 저장은 UTF-8 이 아니라 **CP932** 다.
        한국어 Windows 는 CP949 로 떨어진다. 어느 쪽인지 바이트만 보고는
        확실히 가를 수 없어서, **디코딩 결과의 제목 줄이 우리 열 이름과 몇 개
        맞는지**로 고른다. 어차피 뒤에서 제목 줄을 대조하므로 판정 기준이
        같아지고, 잘못 골랐을 때는 애초에 열이 안 붙는다.
        """
        candidates = (
            ("UTF-8 (BOM)", "utf-8-sig"),
            ("UTF-8", "utf-8"),
            ("CP932（日本語Excel）", "cp932"),
            ("CP949（韓国語Windows）", "cp949"),
        )

        best: tuple[int, str, str] | None = None
        for label, codec in candidates:
            try:
                text = raw.decode(codec)
            except (UnicodeDecodeError, LookupError):
                continue

            head = text.splitlines()[:HEADER_SCAN_ROWS]
            score = 0
            for line in head:
                for delimiter in (",", "\t", ";"):
                    hits, _ = self._header_hits(line.split(delimiter))
                    score = max(score, hits)
            if score >= 2:
                return text, label
            if best is None or score > best[0]:
                best = (score, text, label)

        if best is None:
            raise SheetError(
                "ファイルの文字コードを判別できません。Excelで"
                "「CSV UTF-8（カンマ区切り）」で再度保存してください。"
            )
        return best[1], best[2]

    @staticmethod
    def _sniff_delimiter(text: str) -> str:
        sample = "\n".join(text.splitlines()[:HEADER_SCAN_ROWS])
        try:
            return csv.Sniffer().sniff(sample, delimiters=",\t;").delimiter
        except csv.Error:
            # 열이 1개뿐인 파일에서도 실패한다. 그때는 쉼표로 보아도 결과가 같다.
            return ","

    def _parse_csv(self, raw: bytes, result: SheetImport) -> None:
        text, encoding = self._decode_csv(raw)
        result.kind = "CSV"
        result.encoding = encoding
        reader = csv.reader(
            io.StringIO(text, newline=""), delimiter=self._sniff_delimiter(text)
        )
        self._build_rows([list(cells) for cells in reader], result)

    def _parse_xlsx(self, raw: bytes, result: SheetImport) -> None:
        _Workbook, load_workbook = _openpyxl()
        result.kind = "XLSX"

        try:
            book = load_workbook(
                io.BytesIO(raw),
                read_only=True,
                data_only=True,  # 수식은 계산하지 않고 저장된 값만 읽는다
                keep_links=False,
            )
        except Exception as exc:  # openpyxl 은 형식 오류를 여러 예외로 던진다
            raise SheetError(
                "Excel ファイルを開けません。ファイルが破損していないか、, "
                ".xlsx 形式であるか確認してください。"
            ) from exc

        try:
            sheets = list(book.sheetnames)
            if not sheets:
                raise SheetError("Excel ファイルにシートがありません。")

            # 시트를 순서대로 보며 **제목 줄이 있는 첫 시트**를 쓴다. 표지 시트가
            # 앞에 있는 서식, 우리가 내보낸 「입력 안내」가 섞인 파일 모두를
            # 사람이 시트를 고르지 않고도 통과시키기 위한 것이다.
            matrices: list[tuple[str, list[list[Any]]]] = []
            for name in sheets:
                worksheet = book[name]
                matrix: list[list[Any]] = []
                for cells in worksheet.iter_rows(values_only=True):
                    matrix.append(list(cells))
                    # 상한 + 제목 줄 + 여유. 무한히 큰 시트를 다 읽지 않는다.
                    if len(matrix) > self.max_rows + HEADER_SCAN_ROWS + 2:
                        break
                matrices.append((name, matrix))
                if self._find_header(matrix)[0] >= 0:
                    result.sheet_name = name
                    self._build_rows(matrix, result)
                    return

            # 제목 줄을 못 찾았으면 첫 시트로 위치 대응을 시도한다.
            name, matrix = matrices[0]
            result.sheet_name = name
            self._build_rows(matrix, result)
        finally:
            book.close()

    def parse_upload(self, file_name: str, raw: bytes) -> SheetImport:
        """업로드된 CSV / Excel 을 표의 행으로 바꾼다."""
        if not raw:
            raise SheetError("ファイルが空です。")
        if len(raw) > MAX_UPLOAD_BYTES:
            raise SheetError(
                f"ファイルが大きすぎます。{MAX_UPLOAD_BYTES // (1024 * 1024)}MB 以下に"
                "縮小してください。"
            )

        name = (file_name or "").strip()
        lowered = name.casefold()
        result = SheetImport()

        if lowered.endswith((".csv", ".txt", ".tsv")):
            self._parse_csv(raw, result)
        elif lowered.endswith((".xlsx", ".xlsm")):
            self._parse_xlsx(raw, result)
        elif lowered.endswith(".xls"):
            raise SheetError(
                "旧形式のExcel(.xls)は読み込めません。Excelで"
                "「名前を付けて保存 → Excel ブック(.xlsx)」として保存してください。"
            )
        elif raw[:4] == b"PK\x03\x04":
            # 확장자가 없거나 엉뚱해도 ZIP 머리이면 xlsx 로 보고 시도한다.
            self._parse_xlsx(raw, result)
        else:
            raise SheetError("CSV(.csv) またはExcel(.xlsx)ファイルのみ取り込めます。")

        return result
